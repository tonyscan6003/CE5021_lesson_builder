"""Build the Q&A slide decks (.pptx) from content/*.yaml and the figure engines.

    python Q_A_sessions/build/build_decks.py              # every deck in content/
    python Q_A_sessions/build/build_decks.py --only Q_A_1

Each content file names its engines (modules under Q_A_sessions/engines, with
the lessons' engine contract: FIGURES, build(ctx) -> list[Figure]). The
figures are written to output/<deck id>/figures/ and the deck to
output/<output>. A figure's caption, which states the measured values, is added
to the speaker notes of the slide that shows it.

Slide layouts (the `layout:` of each slide):
    title               dark title slide: title, subtitle, points, aside, figure
    figure_cards        title, a full-width figure, a row of cards (head + text)
    code_figure_below   title, code with side notes, a full-width figure below
    code_figure         title, code and side notes on the left, figure on the right
In code, a line ending in "@batch" or "@epoch" is highlighted in that colour, and
a side note starting with "@batch " or "@epoch " has its lead-in coloured.
"""
from __future__ import annotations

import argparse
import importlib
import sys
from pathlib import Path

import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
QA = ROOT / "Q_A_sessions"
sys.path.insert(0, str(ROOT))

from lessons.engines._common import BuildContext  # noqa: E402

from theme import THEME as T  # noqa: E402

W, H = 13.333, 7.5
M = 0.6                                       # side margin
CONTENT_W = W - 2 * M


def rgb(hex_):
    return RGBColor.from_string(hex_)


# ---------------------------------------------------------------------------
# text helpers

def _font(run, size, color, bold=False, name=None, italic=False):
    f = run.font
    f.size = Pt(size); f.bold = bold; f.italic = italic
    f.name = name or T["body_font"]
    f.color.rgb = rgb(color)


def _bullet(paragraph, char="•", indent=0.22):
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(indent))))
    pPr.set("indent", str(-int(Inches(indent))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(bu)


def textbox(slide, x, y, w, h, *, anchor=MSO_ANCHOR.TOP, margin=0.0, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(margin))
    if name:
        tb.name = name
    return tf


def add_runs(paragraph, parts):
    """parts: list of (text, size, color, bold[, font])"""
    for part in parts:
        text, size, color, bold, *rest = part
        r = paragraph.add_run()
        r.text = text
        _font(r, size, color, bold, rest[0] if rest else None)


def tagged(text):
    """Split a leading '@batch ' / '@epoch ' tag off a side note."""
    for tag in ("batch", "epoch"):
        if text.startswith(f"@{tag} "):
            return tag, text[len(tag) + 2:]
    return None, text


def bullets(tf, items, size=T["body"], color=T["ink"], space=8, first=True):
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if (i == 0 and first) else tf.add_paragraph()
        _bullet(p)
        p.space_after = Pt(space)
        tag, text = tagged(item)
        if tag and ":" in text:
            lead, rest = text.split(":", 1)
            add_runs(p, [(lead + ":", size, T[tag], True), (rest, size, color, False)])
        else:
            add_runs(p, [(text, size, color, False)])


def card(slide, x, y, w, h, fill, name=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w),
                                Inches(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    sh.line.fill.background()
    sh.shadow.inherit = False
    if name:
        sh.name = name
    return sh


def picture(slide, path, x, y, max_w, max_h, center=True):
    iw, ih = Image.open(path).size
    s = min(max_w / iw, max_h / ih)
    w, h = iw * s, ih * s
    if center:
        x = x + (max_w - w) / 2
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    return x, y, w, h


# ---------------------------------------------------------------------------
# slide furniture

def title(slide, text):
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(M), Inches(0.35), Inches(CONTENT_W), Inches(0.8)
    tf = t.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, f"margin_{side}", Inches(0))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    p.text = ""
    add_runs(p, [(text, T["title"], T["ink"], True, T["head_font"])])


def footer(slide, deck, n, dark=False):
    col = T["muted_dark"] if dark else T["muted"]
    tf = textbox(slide, M, H - 0.45, 8, 0.3, anchor=MSO_ANCHOR.MIDDLE, name="Footer")
    add_runs(tf.paragraphs[0], [(deck["footer"], T["small"], col, False)])
    tf = textbox(slide, W - M - 1, H - 0.45, 1, 0.3, anchor=MSO_ANCHOR.MIDDLE, name="Slide number")
    tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
    add_runs(tf.paragraphs[0], [(str(n), T["small"], col, False)])


def notes(slide, text, fig=None):
    body = " ".join((text or "").split())
    if fig is not None:
        body += "\n\nMeasured (figure): " + " ".join(fig.caption.replace("**", "").split())
    slide.notes_slide.notes_text_frame.text = body


def code_box(slide, x, y, w, code, size=T["code"]):
    lines = code.rstrip("\n").split("\n")
    h = 0.3 + len(lines) * size * 1.22 / 72
    card(slide, x, y, w, h, T["bg_alt"], name="Code")
    tf = textbox(slide, x + 0.2, y + 0.15, w - 0.3, h - 0.3, name="Code text")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.0
        tag = None
        for t_ in ("batch", "epoch"):
            if line.rstrip().endswith(f"@{t_}"):
                tag, line = t_, line.rstrip()[: -len(t_) - 1].rstrip()
        add_runs(p, [(line if line else " ", size, T[tag] if tag else T["ink"], bool(tag),
                      T["code_font"])])
    return h


# ---------------------------------------------------------------------------
# layouts

def slide_title(prs, deck, s, figs, n):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    sl.background.fill.solid(); sl.background.fill.fore_color.rgb = rgb(T["dark"])
    tf = textbox(sl, 0.8, 1.2, 7.4, 1.2, anchor=MSO_ANCHOR.BOTTOM, name="Title")
    add_runs(tf.paragraphs[0], [(deck["title"], 44, T["white"], True, T["head_font"])])
    tf = textbox(sl, 0.8, 2.45, 7.4, 0.5, name="Subtitle")
    add_runs(tf.paragraphs[0], [(deck["subtitle"], 20, T["accent_light"], False)])
    tf = textbox(sl, 0.8, 3.3, 7.2, 2.2, name="Contents")
    bullets(tf, s["points"], size=18, color=T["white"], space=6)
    tf = textbox(sl, 0.8, 5.55, 7.0, 1.1, name="Note")
    add_runs(tf.paragraphs[0], [(" ".join(s["aside"].split()), T["small"] + 2, T["muted_dark"], False)])
    if s.get("figure"):
        card(sl, 8.75, 1.3, 3.9, 3.9, T["white"], name="Figure card")
        picture(sl, figs[s["figure"]].path, 8.9, 1.45, 3.6, 3.6)
        tf = textbox(sl, 8.75, 5.3, 3.9, 0.6, name="Figure label")
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        add_runs(tf.paragraphs[0], [("the toy problem: three classes", T["small"] + 1,
                                     T["muted_dark"], False)])
    footer(sl, deck, n, dark=True)
    notes(sl, s.get("notes"), figs.get(s.get("figure")))


def slide_figure_cards(prs, deck, s, figs, n):
    sl = prs.slides.add_slide(prs.slide_layouts[5])
    title(sl, s["title"])
    f = figs[s["figure"]]
    refs = s.get("refs")
    x, y, w, h = picture(sl, f.path, M, 1.25, CONTENT_W, 3.15 if refs else 3.35)
    cards = s["cards"]
    gap = 0.3
    cw = (CONTENT_W - gap * (len(cards) - 1)) / len(cards)
    cy = y + h + 0.25
    ch = min(1.55, (H - 0.55 - (0.5 if refs else 0.1)) - cy)
    for i, c in enumerate(cards):
        cx = M + i * (cw + gap)
        card(sl, cx, cy, cw, ch, T["accent_soft"], name=f"Card {i + 1}")
        tf = textbox(sl, cx + 0.2, cy + 0.14, cw - 0.4, ch - 0.24, name=f"Card {i + 1} text")
        add_runs(tf.paragraphs[0], [(c["head"], T["body"] + 2, T["accent_dark"], True)])
        p = tf.add_paragraph(); p.space_before = Pt(3)
        add_runs(p, [(" ".join(c["text"].split()), T["body"], T["ink"], False)])
    if s.get("refs"):
        tf = textbox(sl, M, cy + ch + 0.08, CONTENT_W, 0.4, name="References")
        add_runs(tf.paragraphs[0], [(" ".join(s["refs"].split()), T["small"], T["muted"], False)])
    footer(sl, deck, n)
    notes(sl, s.get("notes"), f)


def slide_code_figure_below(prs, deck, s, figs, n):
    sl = prs.slides.add_slide(prs.slide_layouts[5])
    title(sl, s["title"])
    cw = 6.3
    h = code_box(sl, M, 1.3, cw, s["code"])
    tf = textbox(sl, M + cw + 0.4, 1.3, CONTENT_W - cw - 0.4, max(h, 1.5), name="Notes")
    bullets(tf, s["side"], size=T["body"], space=4)
    f = figs[s["figure"]]
    picture(sl, f.path, M, 1.3 + max(h, 1.5) + 0.25, CONTENT_W, H - 0.55 - (1.3 + max(h, 1.5) + 0.25))
    footer(sl, deck, n)
    notes(sl, s.get("notes"), f)


def slide_code_figure(prs, deck, s, figs, n):
    sl = prs.slides.add_slide(prs.slide_layouts[5])
    title(sl, s["title"])
    cw = 6.75
    h = code_box(sl, M, 1.3, cw, s["code"])
    tf = textbox(sl, M, 1.3 + h + 0.25, cw, H - 0.6 - (1.3 + h + 0.25), name="Notes")
    bullets(tf, s["side"], size=T["body"], space=5)
    f = figs[s["figure"]]
    fx = M + cw + 0.35
    picture(sl, f.path, fx, 1.3, W - M - fx, H - 0.6 - 1.3)
    footer(sl, deck, n)
    notes(sl, s.get("notes"), f)


LAYOUTS = {"title": slide_title, "figure_cards": slide_figure_cards,
           "code_figure_below": slide_code_figure_below, "code_figure": slide_code_figure}


# ---------------------------------------------------------------------------

def build_deck(path: Path) -> Path:
    deck = yaml.safe_load(path.read_text())
    out_dir = QA / "output" / deck["id"]
    ctx = BuildContext(out_dir, out_dir / "figures")
    figs = {}
    for name in deck["engines"]:
        mod = importlib.import_module(f"Q_A_sessions.engines.{name}")
        print(f"  engine {name}:", end=" ", flush=True)
        for f in mod.build(ctx):
            f.path = out_dir / f.png
            figs[f.id] = f
            print(f.id, end=" ", flush=True)
        print()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = deck["title"]
    prs.core_properties.subject = deck["subtitle"]
    for n, s in enumerate(deck["slides"], 1):
        LAYOUTS[s["layout"]](prs, deck, s, figs, n)
    out = QA / "output" / deck["output"]
    prs.save(out)
    print(f"  -> {out.relative_to(ROOT)} ({len(deck['slides'])} slides)")
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--only", help="deck id, e.g. Q_A_1")
    a = ap.parse_args()
    for path in sorted((QA / "content").glob("*.yaml")):
        if a.only and not path.stem.startswith(a.only):
            continue
        print(path.name)
        build_deck(path)


if __name__ == "__main__":
    main()
